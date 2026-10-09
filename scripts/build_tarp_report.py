"""TARP report (VIT Chennai template) for StanceSense-RT -> docs/tarp_report/.

Follows the template's page specification (A4; left 3.81 cm, other margins 2.54 cm;
Times New Roman) and its order: title page, declaration, certificate, abstract,
contents, lists of figures / tables / acronyms, six chapters, APA references.
Tables and result numbers are read from the saved artefacts at build time:
  models/rule_classifier/*.json|csv     rule-based squat classifier
  models/experiments/multiview_A.json   class balance of the cleaned data (share of non-squat windows)
  config/stance_rules.yaml              stance rule constants
  models/pose_speed.json                timing of the pretrained pose model on this laptop
  data/stancesense.db (session 10)      the live assessment in Section 5.6
  demo_output/end_to_end_*.txt          rotation arcs of the offline MM-Fit demonstration
and the number of automated tests is counted with pytest at build time.
Figures come from scripts/make_tarp_figures.py (diagrams and results) and
scripts/capture_report_media.py (recording stills, dashboard screenshots). After
building, Word updates the contents / figure / table lists (scripts/tarp_finalize.ps1).
    .venv/Scripts/python.exe scripts/make_tarp_figures.py
    .venv/Scripts/python.exe scripts/capture_report_media.py
    .venv/Scripts/python.exe scripts/build_tarp_report.py
"""
from __future__ import annotations

import csv
import glob
import json
import os
import re
import subprocess
import sys
from collections import Counter
from datetime import datetime

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
os.chdir(ROOT)
OUTDIR = os.path.join("docs", "tarp_report")
FIG = os.path.join(OUTDIR, "figures")
OUT = os.path.join(OUTDIR, "TARP_Report_StanceSense-RT.docx")
if len(sys.argv) > 1:                      # optional: build somewhere else (e.g. to check it first)
    OUT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", sys.argv[1]))         if not os.path.isabs(sys.argv[1]) else sys.argv[1]

# Front matter as the candidate filled it in (title page, declaration, certificate).
CANDIDATE, REG_NO = "Shakthisurya", "23MIA1151"
GUIDE = "Aravindkumar S"
SIGN_DATE = "07-10-2026"

TITLE = "Squat Form Analysis using Integrated Markerless Hip-Rotation Profiling and Real-Time Pose Estimation"
DEGREE = "M.Tech. (Integrated) in Computer Science and Engineering with Specialization in Business Analytics"
FONT = "Times New Roman"


def J(p):
    return json.load(open(p))


RC = J("models/rule_classifier/metrics_eval.json")
RCD = J("models/rule_classifier/metrics_dev.json")
EVAL_ROWS = list(csv.DictReader(open("models/rule_classifier/results_eval.csv")))
E2E = J("models/rule_classifier/e2e_realtime.json")
EXP = {"A": J("models/experiments/multiview_A.json")}
RULES = yaml.safe_load(open("config/stance_rules.yaml"))
M = RC["metrics"]
_front = [r for r in EVAL_ROWS if r["view"] == "front"]
FRONT_ACC = sum(int(r["correct"]) for r in _front) / len(_front)
# share of non-squat 6-second windows in the cleaned data (the class imbalance behind the rule-based design)
MAJORITY = EXP["A"]["data"]["majority_rate"]

# pretrained pose model: file sizes and timing on this laptop (scripts/capture_report_media.py)
SPEED = J("models/pose_speed.json")
CPU = re.sub(r"\((R|TM)\)", "", SPEED["cpu"]).replace("  ", " ").strip()

# the live assessment shown in Section 5.6, read back from the application's database
sys.path.insert(0, "src")
from stancesense.storage import Store  # noqa: E402
from stancesense.squat.rule_classifier import FRONT_FACING_MAX_DEG  # noqa: E402
LIVE_ID = 10
with Store(os.path.join("data", "stancesense.db")) as _st:
    LIVE_P, LIVE_S = _st.profile_for(LIVE_ID), _st.stance_for(LIVE_ID)
    LIVE_SUM, LIVE_SQ = _st.rep_summary(LIVE_ID), _st.squat_run_for(LIVE_ID)
    LIVE_DATE = datetime.fromisoformat(next(x["created_at"] for x in _st.sessions() if x["id"] == LIVE_ID))
LIVE_DAY = f"{LIVE_DATE.day} {LIVE_DATE:%B %Y}"

# largest total rotation arc printed by the offline MM-Fit demonstration
DEMO_ARCS, DEMO_TOTAL = {}, {}
for _f in glob.glob(os.path.join("demo_output", "end_to_end_*.txt")):
    _txt, _wid = open(_f).read(), os.path.basename(_f)[len("end_to_end_"):-len(".txt")]
    _m = re.search(r"arc ([0-9.]+) deg", _txt)
    if _m:
        DEMO_ARCS[_wid] = float(_m.group(1))
    _m = re.search(r"TOTAL true (\d+), counted (\d+), confirmed (\d+)", _txt)
    if _m:
        DEMO_TOTAL[_wid] = tuple(int(g) for g in _m.groups())       # (true, counted, confirmed)
DEMO_WID, DEMO_ARC = max(DEMO_ARCS.items(), key=lambda kv: kv[1])

# number of automated tests, counted rather than typed
_col = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q", "-p", "no:cacheprovider"],
                      capture_output=True, text=True)
N_TESTS = int(re.search(r"(\d+) tests? collected", _col.stdout).group(1))


def pc(x, d=1):
    return f"{x * 100:.{d}f}"


# ===================================================================== document
doc = Document()


def set_font(style, size, bold=None, italic=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = color or RGBColor(0, 0, 0)


normal = doc.styles["Normal"]
set_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

h1 = doc.styles["Heading 1"]; set_font(h1, 14, bold=True)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(0); h1.paragraph_format.space_after = Pt(18)
h1.paragraph_format.line_spacing = 1.5; h1.paragraph_format.page_break_before = True
h2 = doc.styles["Heading 2"]; set_font(h2, 12, bold=True)
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(6); h2.paragraph_format.keep_with_next = True
h3 = doc.styles["Heading 3"]; set_font(h3, 12, bold=True, italic=True)
h3.paragraph_format.space_before = Pt(8); h3.paragraph_format.space_after = Pt(4); h3.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 12)
    doc.styles[name].paragraph_format.line_spacing = 1.5
    doc.styles[name].paragraph_format.space_after = Pt(3)
for name in ("Figure Caption", "Table Caption"):
    st = doc.styles.add_style(name, 1)
    st.base_style = normal
    set_font(st, 11)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_before = Pt(3); st.paragraph_format.space_after = Pt(12 if name == "Figure Caption" else 4)
doc.styles["Table Caption"].paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3.81), Cm(2.54)
sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
TEXT_W = 21.0 - 3.81 - 2.54


def para(text="", size=None, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=None,
         spacing=None, keep=False, highlight=False, underline=False):
    p = doc.add_paragraph()
    p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if spacing is not None:
        p.paragraph_format.line_spacing = spacing
    if keep:
        p.paragraph_format.keep_with_next = True
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.underline = underline
        if size:
            r.font.size = Pt(size)
        if highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def runs(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, after=None, spacing=None, indent=None):
    """Paragraph from [(text, {'b':..,'i':..,'h':..})] pieces."""
    p = doc.add_paragraph(); p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if spacing is not None:
        p.paragraph_format.line_spacing = spacing
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("b", False); r.italic = f.get("i", False)
        if size:
            r.font.size = Pt(size)
        if f.get("h"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def P(text):
    return para(text)


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(it)


def H2(text):
    return doc.add_heading(text, level=2)


def H3(text):
    return doc.add_heading(text, level=3)


def chapter(n, title):
    p = doc.add_heading("", level=1)
    p.add_run(f"CHAPTER {n}")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run(title.upper())
    return p


def front_heading(text):
    """Front-matter heading that appears in the contents (LIST OF ...)."""
    p = doc.add_heading(text, level=1)
    return p


FIGS, TABS = [], []


def figure(fname, width_cm, caption, number):
    doc.add_picture(os.path.join(FIG, fname), width=Cm(width_cm))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(style="Figure Caption")
    r = p.add_run(f"Figure {number}: "); r.bold = True
    p.add_run(caption)
    FIGS.append(number)


def shade(cell, hexcolor="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)


def table(number, caption, headers, rows, widths, size=10.5, align_first_left=True, left_cols=()):
    p = doc.add_paragraph(style="Table Caption")
    r = p.add_run(f"Table {number}: "); r.bold = True
    p.add_run(caption)
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.width = Cm(widths[j]); shade(c)
        cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.0; cp.paragraph_format.space_after = Pt(0)
        rr = cp.add_run(h); rr.bold = True; rr.font.size = Pt(size)
    trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].width = Cm(widths[j])
            cp = cells[j].paragraphs[0]
            cp.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and align_first_left) or j in left_cols
                            else WD_ALIGN_PARAGRAPH.CENTER)
            cp.paragraph_format.line_spacing = 1.0; cp.paragraph_format.space_after = Pt(0)
            rr = cp.add_run(str(v)); rr.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    TABS.append(number)
    return t


def equation(text, number):
    p = doc.add_paragraph()
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("\t" + text); r.italic = True
    p.add_run(f"\t({number})")


def add_field(p, instr, placeholder=" "):
    r = p.add_run(); b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    r = p.add_run(); s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r._r.append(s)
    p.add_run(placeholder)
    r = p.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)


def page_number_format(section, fmt):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt); pg.set(qn("w:start"), "1")


def footer_page_number(section):
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", "1")


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def logo(width_cm=8.0):
    doc.add_picture(os.path.join(FIG, "vit_logo_1.png"), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ===================================================================== front matter
C = WD_ALIGN_PARAGRAPH.CENTER
logo(8.5)
para("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING", size=16, bold=True, align=C, after=4)
para("October, 2026", size=14, align=C)
para("")
para("A project report on", size=12, italic=True, align=C, after=6)
para(TITLE.upper(), size=16, bold=True, align=C, after=18, spacing=1.2)
para("Submitted in partial fulfillment for the award of the degree of", size=14, italic=True, align=C, after=10)
para(DEGREE, size=20, bold=True, align=C, after=30, spacing=1.15)
para("by", size=14, italic=True, align=C, after=10)
para(f"{CANDIDATE} ({REG_NO})", size=16, bold=True, align=C, after=2)

logo(8.0)
doc.paragraphs[-1].paragraph_format.page_break_before = True
para("DECLARATION", size=14, bold=True, align=C, after=24, underline=False)
runs([("I hereby declare that the thesis entitled “", {}), (TITLE, {"b": True}), ("” submitted by ", {}),
      (f"{CANDIDATE} ({REG_NO})", {}),
      (f", for the award of the degree of {DEGREE}, Vellore Institute of Technology, Chennai is a record of bonafide "
       "work carried out by me under the supervision of ", {}), (GUIDE, {}), (".", {})],
     size=13, indent=1.27, after=12)
runs([("I further declare that the work reported in this thesis has not been submitted and will not be submitted, "
       "either in part or in full, for the award of any other degree or diploma in this institute or any other "
       "institute or university.", {})], size=13, indent=1.27, after=60)
para("Place: Chennai", size=13, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
p = para("", align=WD_ALIGN_PARAGRAPH.LEFT)
p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
r = p.add_run(f"Date: {SIGN_DATE}\tSignature of the Candidate"); r.font.size = Pt(13)
page_break()

logo(8.0)
para("School of Computer Science and Engineering", size=14, bold=True, align=C, after=12)
para("CERTIFICATE", size=16, align=C, after=18)
runs([("This is to certify that the report entitled “", {}), (TITLE, {"b": True}),
      ("” is prepared and submitted by ", {}), (CANDIDATE, {}),
      (f" to Vellore Institute of Technology, Chennai, in partial fulfillment of the requirement for the award of "
       f"the degree of {DEGREE} is a bonafide record carried out under my guidance. The project fulfills the "
       "requirements as per the regulations of this University and in my opinion meets the necessary standards for "
       "submission. The contents of this report have not been submitted and will not be submitted either in part or "
       "in full, for the award of any other degree or diploma and the same is certified.", {})],
     size=12, after=48)
para("Signature of the Faculty:", align=WD_ALIGN_PARAGRAPH.LEFT, after=18)
runs([("Name: Dr./Prof. ", {}), (GUIDE, {})], align=WD_ALIGN_PARAGRAPH.LEFT, after=18)
para(f"Date: {SIGN_DATE}", align=WD_ALIGN_PARAGRAPH.LEFT)

# ---- section 2: roman numbering
s2 = doc.add_section(WD_SECTION.NEW_PAGE)
page_number_format(s2, "lowerRoman"); footer_page_number(s2)
doc.sections[0].footer.is_linked_to_previous = False

para("ABSTRACT", size=14, bold=True, align=C, after=18, underline=True)
ABSTRACT = [
    "Squat coaching usually starts from one rule: feet about shoulder-width apart, toes turned out a little. Hips "
    "differ, though. How far each hip turns inward and outward varies from person to person, and that range shapes "
    "how wide and how turned-out a comfortable stance is. Clinicians check hip rotation with a goniometer or with "
    "Craig's test, and both need a trained examiner. This project runs a functional version of that check with one "
    "webcam and then watches the person squat.",
    "StanceSense-RT runs the pretrained MediaPipe BlazePose model, unchanged, on a laptop CPU. The user first sits and "
    "turns each lower leg to its inward and outward limits. The shank's angle from the vertical gives a functional "
    "rotation reading, and explicit rules turn the readings into a starting stance width and toe-out angle. The user "
    "then squats. A rule-based classifier with no trained parameters follows each movement through standing, "
    "descent, bottom and ascent, and accepts it as a squat only when it passes at least seven of nine biomechanical "
    "checks. The confirmed repetitions are measured for depth, trunk lean, stance width and knee tracking, which gives "
    "a narrow, moderate or wide stance suggestion. Results are stored in SQLite and shown in a Streamlit dashboard.",
    f"The classifier was tested on a public multi-view fitness video dataset after 47 duplicate recordings were "
    f"removed. Its thresholds were set on {len(RC['development_people'])} participants and then frozen. On "
    f"{M['videos']} videos from {len(RC['evaluation_people'])} other participants it reached {pc(M['accuracy'])}% "
    f"accuracy, {pc(M['precision'])}% precision, {pc(M['recall'])}% recall, an F1-score of {pc(M['f1'])}% and "
    f"{pc(M['specificity'])}% specificity, with {pc(FRONT_ACC)}% accuracy on front-camera videos, the setting the "
    "application uses. Rules were preferred to a trained classifier because the labelled data is small and mostly "
    "non-squat, and because a rule can say why a repetition was rejected. The rotation readings are reported "
    "qualitatively, since they were not checked against a goniometer.",
]
for t in ABSTRACT:
    para(t, spacing=1.3)          # keeps the abstract on one page
runs([("Keywords: ", {"b": True}), ("squat analysis, markerless pose estimation, MediaPipe BlazePose, hip rotation, "
                                     "stance recommendation, rule-based classification, repetition counting", {"i": True})],
     spacing=1.3)

toc_title = para("TABLE OF CONTENTS", size=14, bold=True, align=C, after=12)
toc_title.paragraph_format.page_break_before = True   # no empty page when the abstract fills its page
p = doc.add_paragraph(); add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field.")

front_heading("LIST OF FIGURES")
p = doc.add_paragraph(); add_field(p, 'TOC \\h \\z \\t "Figure Caption,1"', "Right-click and choose Update Field.")
front_heading("LIST OF TABLES")
p = doc.add_paragraph(); add_field(p, 'TOC \\h \\z \\t "Table Caption,1"', "Right-click and choose Update Field.")

front_heading("LIST OF ACRONYMS")
ACR = [("CPU", "Central Processing Unit"),
       ("CT", "Computed Tomography"), ("ER", "External Rotation"),
       ("F1", "Harmonic mean of precision and recall"), ("FN / FP", "False Negative / False Positive"),
       ("FNR / FPR", "False Negative Rate / False Positive Rate"), ("fps", "Frames Per Second"),
       ("GPU", "Graphics Processing Unit"), ("IR", "Internal Rotation"),
       ("MRI", "Magnetic Resonance Imaging"),
       ("RGB / RGB-D", "Colour video / colour plus depth video"), ("RT", "Real Time"),
       ("TN / TP", "True Negative / True Positive"), ("UI", "User Interface")]
t = doc.add_table(rows=0, cols=2)
for a, b in ACR:
    cells = t.add_row().cells
    cells[0].width, cells[1].width = Cm(3.5), Cm(11.0)
    for c, v, bold in ((cells[0], a, True), (cells[1], b, False)):
        cp = c.paragraphs[0]; cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.line_spacing = 1.3
        rr = cp.add_run(v); rr.bold = bold

# ---- section 3: arabic numbering
s3 = doc.add_section(WD_SECTION.NEW_PAGE)
page_number_format(s3, "decimal")

# ===================================================================== CHAPTER 1
chapter(1, "Introduction")
H2("1.1 Background")
P("The squat loads the hips, knees and ankles together and appears in nearly every strength and rehabilitation "
  "programme, from beginner routines to competitive powerlifting. It is usually taught with the same generic cue: "
  "feet about shoulder-width apart, toes turned out a little, sit down between the heels.")
P("That cue assumes everyone's hips move alike, and they do not. How far a hip rotates inward (internal rotation, "
  "IR) and outward (external rotation, ER) depends on the shape of the femur and the hip socket and on soft-tissue "
  "flexibility. People whose hips turn outward more than inward tend to prefer a wider, more turned-out stance; the "
  "opposite pattern often suits a narrower, straighter one. Hip rotation is measured with a goniometer, and femoral "
  "version is estimated with Craig's test (Ruwe et al., 1992), but both need a trained examiner, so most people who "
  "squat at home or in a gym never have either done.")
P("Pose estimation, meanwhile, has become cheap. MediaPipe BlazePose finds 33 body landmarks from a single colour "
  "camera in real time on a laptop CPU (Bazarevsky et al., 2020), so joint angles and whole movements can now be "
  "measured with a webcam.")
H2("1.2 Motivation")
P("This project asks how far one webcam can go in personalising squat advice. Can a seated rotation test in front "
  "of a camera show a person their own hip-rotation pattern? And can the same camera then confirm that the person is "
  "really squatting, count the repetitions and judge whether the stance suits them?")
P(f"Trust in the results mattered just as much. Labelled squat data is scarce and unbalanced: in the public dataset "
  f"we used, {pc(MAJORITY)}% of all six-second stretches of video show an exercise other than a squat, so a "
  f"classifier can score high accuracy while recognising few squats. We therefore made the squat decision with "
  f"explicit biomechanical rules whose thresholds can be read and checked, set them on a few development "
  f"participants, and report every result on other people.")
H2("1.3 Problem Statement")
P("From live video of one person and a single webcam, the system has to measure a functional hip-rotation profile "
  "for each leg and turn it into a starting stance with rules a coach can inspect. During the squats it has to decide "
  "in real time whether each movement is a squat, count and measure the confirmed repetitions, and suggest keeping, "
  "narrowing or widening the stance. It must run on a laptop CPU and be tested on people not used to design it.")
H2("1.4 Objectives")
bullets(["A guided, camera-only seated protocol that measures functional hip IR and ER, locking each reading when "
         "the user holds still.",
         "A hip profile from both legs, mapped to a starting stance by transparent rules.",
         "A rule-based squat classifier and repetition counter that works from pose landmarks alone.",
         "A narrow, moderate or wide stance suggestion from the squats the person performs.",
         "Stored assessments and a dashboard that runs locally, in Docker or as a desktop app.",
         "One evaluation of the frozen classifier on unseen participants, plus tests of the live pipeline on raw "
         "video."], style="List Number")
H2("1.5 Scope and Limitations")
P("The rotation measurement is functional: it reads how far the lower leg swings when the hip rotates, and it is "
  "not a measure of femoral or acetabular version, which needs CT or MRI. The readings were not validated against a "
  "goniometer, so no claim is made about their angular accuracy. The squat check expects one person facing the "
  "camera, and it reports the measurements behind each decision without diagnosing named form faults.")
H2("1.6 Organisation of the Report")
P("Chapter 2 reviews pose estimation, squat biomechanics, clinical hip-rotation tests and video-based exercise "
  "assessment. Chapter 3 describes the system, Chapter 4 the data, metrics and implementation, and Chapter 5 the "
  "results. Chapter 6 concludes and lists future work.")

# ===================================================================== CHAPTER 2
chapter(2, "Related Work")
H2("2.1 Markerless Human Pose Estimation")
P("Pose estimation locates body joints in images. OpenPose (Cao et al., 2021) made real-time multi-person 2D pose "
  "practical on a GPU by linking keypoints into skeletons with part affinity fields. BlazePose (Bazarevsky et al., "
  "2020) targets one person, on the device, in real time: a detector finds the person and a tracking network predicts "
  "33 landmarks per frame, including the hands and feet. Through MediaPipe (Lugaresi et al., 2019) it also returns "
  "world landmarks in metres, centred between the hips, with a visibility score for each. StanceSense-RT measures "
  "everything from the world landmarks and uses the visibility scores to ignore joints the network is guessing. "
  "Depth is the weakest axis from one camera; the same squats read as shallower and more forward-leaning from the "
  "front than from the side (Section 5.4), so rules have to be set for the camera position in use.")
H2("2.2 Squat Biomechanics and Stance Width")
P("Escamilla (2001) reviewed knee loads during the dynamic squat and how they change with depth and technique. "
  "Escamilla et al. (2001) filmed lifters at narrow, medium and wide stances and found that stance width changes the "
  "hip and knee angles and the joint moments. Schoenfeld (2010) reviewed how depth, stance width, foot position and "
  "trunk inclination alter muscle activity and joint loading. This project rests on two ideas from that work: stance "
  "width and toe-out are adjustable rather than fixed, and depth and trunk lean are measurable signs of how well a "
  "person copes with a given stance.")
H2("2.3 Clinical Measurement of Hip Rotation")
P("Hip rotation is normally measured with a goniometer, with the person lying face down or seated with the knee at "
  "90 degrees; the lower leg is the pointer and its angle from the vertical is the rotation. Roach and Miles (1991) "
  "reported normal active hip and knee ranges from a large population survey and how they change with age. Femoral "
  "version, the twist of the thigh bone that partly sets the rotation range, is measured with CT or MRI and estimated "
  "clinically with Craig's test, which Ruwe et al. (1992) compared with imaging and intra-operative measurements. "
  "Our seated protocol uses the same geometry, with the camera in place of the goniometer, and gives a functional "
  "screen rather than a measurement of bone shape.")
H2("2.4 Exercise Recognition, Counting and Form Assessment")
P("Repetition counting from video has mostly relied on learned models. RepNet (Dwibedi et al., 2020) counts "
  "repetitions of arbitrary actions from the periodicity of learned frame embeddings. MM-Fit (Strömbäck et al., 2020) "
  "recorded full-body workouts with phones, watches, earbuds and RGB-D cameras, with 2D and 3D pose and labelled "
  "exercise sets, and its authors used it for exercise recognition and counting with multimodal deep networks.")
P("For form assessment, Fitness-AQA (Parmar et al., 2022) collected back squats, rows and overhead presses from "
  "social-media videos with trainer-labelled errors such as knees caving in, and showed that self-supervised "
  "pose-aware features beat off-the-shelf pose estimators. EC3D (Zhao et al., 2022) recorded squats, lunges and "
  "planks with deliberate mistakes and trained a graph network on 3D pose to classify and correct them. For "
  "rehabilitation, KIMORE (Capecci et al., 2019) has Kinect recordings and clinical scores for 78 people doing five "
  "exercises, one a squat, and UI-PRMD (Vakanski et al., 2018) has motion-capture and Kinect data for ten movements, "
  "including the deep squat, done correctly and incorrectly by ten people. Two problems recur: some datasets have "
  "very few people (four in EC3D, ten in UI-PRMD), and others have no participant identities, so one person can sit "
  "in both training and test data. We met a version of the second problem in our own dataset (Section 4.1.1).")
H2("2.5 Research Gap")
P("We found no system that links a person's hip-rotation pattern to a stance recommendation and then checks that "
  "stance during live squatting with one webcam. Squat-form systems mostly ask whether a repetition is good or bad, "
  "and they answer with learned models whose decisions are hard to explain to the person squatting. StanceSense-RT "
  "combines a functional hip screen, explicit stance rules and a transparent squat check that reports each decision "
  "with the measurements behind it.")

# ===================================================================== CHAPTER 3
chapter(3, "Proposed Methodology")
H2("3.1 System Overview")
P("Figure 3.1 shows the pipeline. Each webcam frame goes through MediaPipe BlazePose, which returns 33 world "
  "landmarks with visibility scores. A short calibration measures the person's hip width, femur length and the "
  "resting angle of each shank. In the seated stage the person turns each lower leg to its inward and outward "
  "limits, and the readings become a hip profile and a starting stance. In the standing stage the person squats; a "
  "rule-based classifier confirms each repetition, the confirmed repetitions are measured, and the measurements give "
  "a squat-based stance suggestion. Everything is saved to a local SQLite database and shown in a dashboard.")
figure("fig_architecture.png", 11.5, "StanceSense-RT pipeline. The only trained component in the live path is the "
       "pretrained pose network; the squat decision is made by explicit rules.", "3.1")
P("The camera loop runs as its own process with a native OpenCV window at the camera's full frame rate; the "
  "Streamlit dashboard starts it, follows it through a small JSON status file and reads the result from the "
  "database. The assessment logic sits in a controller with no camera or interface code, so tests can drive it with "
  "recorded landmarks.")
H2("3.2 Pose Estimation and Body-Scale Calibration")
H3("3.2.1 The pretrained pose model")
P("StanceSense-RT trains no pose model of its own. Every frame goes through MediaPipe BlazePose (Bazarevsky et al., "
  "2020), a pretrained network that Google publishes with the MediaPipe framework (Lugaresi et al., 2019) for "
  "real-time tracking of one person on a phone or laptop. It is used as a fixed feature extractor: its weights are "
  "never changed, and everything after it is explicit geometry and thresholds.")
P("BlazePose runs in two stages (Table 3.1). A detector finds the person, using the face as its main cue to where "
  "the body is (Bazarevsky et al., 2020), and a landmark network then predicts 33 landmarks inside that region. Each "
  "landmark has normalised image coordinates, which the application uses for drawing, world coordinates in metres "
  "centred between the hips, which it uses for every measurement, and a visibility score. On the next frame the "
  "region comes from the previous landmarks, so the detector runs again only when tracking is lost.")
_mf = SPEED["model_files_mb"]
table("3.1", f"Pretrained pose model used (MediaPipe {SPEED['mediapipe']}, weights unchanged)",
      ["Network", "File (size)", "Role", "Output"],
      [["Person detector", f"pose_detection.tflite ({_mf['pose_detection.tflite']} MB)",
        "finds the person when tracking starts or is lost", "region around the person"],
       ["Landmark model, full variant (model complexity 1)",
        f"pose_landmark_full.tflite ({_mf['pose_landmark_full.tflite']} MB)",
        "predicts the body landmarks in that region",
        "33 landmarks: image and world (metre) coordinates, visibility"]],
      [3.6, 4.0, 3.6, 3.4], size=10, left_cols=(2, 3))
P(f"On this project's laptop ({CPU}) the two networks took a median of {SPEED['median_ms']} ms per frame (90th "
  f"percentile {SPEED['p90_ms']} ms) over {SPEED['frames_timed']} frames of a front-view squat video scaled to 640 "
  f"pixels, about {SPEED['fps_at_median']:.0f} frames per second, and found the person in "
  f"{pc(SPEED['pose_found_frac'], 0)}% of them. That leaves time for the rest of the pipeline within each frame of a "
  f"30 frames-per-second webcam. The model files ship with the MediaPipe package, so nothing is downloaded at run time.")
P("A pretrained model was the practical choice. Training a pose network takes many thousands of images with "
  "hand-labelled joints and a GPU, and none of the exercise datasets here labels joints. Keeping the learned part to "
  "the pose network also means the application can name the check that failed whenever it rejects a repetition.")
H3("3.2.2 Settings and calibration")
P("MediaPipe Pose runs at model complexity 1 with landmark smoothing on and confidence thresholds of 0.5, and all "
  "measurements use the world landmarks (y points down). A leg is measured only while its hip, knee and ankle all "
  "have visibility of at least 0.6; otherwise the user is asked to move back into frame and the reading is dropped.")
P("Calibration takes 45 frames (about 1.5 seconds at 30 frames per second) while the person sits still. The median "
  "distance between the two hip landmarks becomes the hip width, the median hip-to-knee distance the femur length, "
  "and each leg's median shank angle its neutral offset. Later lengths are divided by the hip width or the femur "
  "length, so the rules do not depend on the person's height or distance from the camera.")
figure("fig_pose_overlay.png", 11.5, "MediaPipe BlazePose landmarks on three frames of a front-view squat from the "
       "Multi-View dataset (Prashanth et al., 2026; CC BY 4.0; faces blurred by the dataset authors).", "3.2")
H2("3.3 Functional Hip-Rotation Measurement")
P("The person sits facing the camera with the knees bent to about 90 degrees and the lower legs hanging free. With "
  "the thigh still, rotating the hip about the femur's long axis swings the shank sideways, and from the front the "
  "angle between the shank (knee to ankle) and the vertical is the rotation reading (Figure 3.3). The shank is long "
  "and well tracked, which makes it a steadier pointer than the hip joint centre, the least reliable landmark in "
  "markerless pose estimation.")
figure("fig_shank_geometry.png", 11.0, "Rotation geometry for the right leg, seen from the front. The dark line is the "
       "thigh, held still, and the dashed line is the vertical. Internal rotation swings the ankle away from the "
       "midline and reads positive; external rotation reads negative.", "3.3")
P("For a knee at (xₖ, yₖ) and an ankle at (xₐ, yₐ), the reading is the signed angle between the shank "
  "and the downward vertical:")
equation("θ = atan2(xₐ − xₖ , yₐ − yₖ) − θₙₑᵤₜ", "3.1")
P("The sign is flipped for the left leg so that internal rotation is positive on both sides, and θ neutral is the "
  "angle locked in the neutral step. Each reading is the median of the last five frames. Every leg goes through three "
  "steps, each opening with a four-second countdown that shows the next instruction while nothing is measured. For "
  "neutral the person points the foot straight ahead, and the reading locks once it stays within 4 degrees for one "
  "second. Then the person turns the leg as far as it goes one way and then the other. A reading counts only when it "
  "is at least 5 degrees from neutral, and it locks after staying within 3 degrees of the furthest point reached in "
  "that step, steady to within 4 degrees, for two seconds; the second extreme must have the opposite sign. The "
  "positive locked value is the leg's internal rotation and the size of the negative one its external rotation. "
  "Holds are timed in seconds, so the protocol behaves the same at any frame rate.")
H2("3.4 Hip Profile and Starting-Stance Rules")
P("The four locked values make up the profile, where l and r mark the left and right legs and s either side:")
equation("IR = (IRₗ + IRᵣ) / 2,   ER = (ERₗ + ERᵣ) / 2,   arc = IR + ER", "3.2")
equation("biasₛ = (IRₛ − ERₛ) / 2", "3.3")
equation("bias = (biasₗ + biasᵣ) / 2,   symmetry = |biasₗ − biasᵣ|", "3.4")
P("A bias above +5 degrees is labelled IR-dominant, below −5 degrees ER-dominant, and anything between balanced. "
  "The starting stance follows from the external-rotation dominance d = −bias. A hip that turns outward more needs "
  "more room to open, so the toe-out and the width both grow with d, and a small total arc damps the width change:")
equation("toe-out = clip(15 + 0.8·d, 5, 30) degrees", "3.5")
equation("width = clip(1.0 + 0.02·d·s, 0.9, 1.6) × hip width", "3.6")
equation("s = clip(arc / 60, 0.5, 1.0)", "3.7")
P("If the two legs' biases differ by more than 8 degrees, the stance is flagged as asymmetric. All constants live in "
  "one configuration file (Table 3.2), so a coach can read and change the rules without touching the code. The "
  "output is a starting stance that the squat check later confirms or adjusts.")
w = RULES["width"]; to = RULES["toe_out"]
table("3.2", "Constants of the starting-stance rules (config/stance_rules.yaml)",
      ["Constant", "Value", "Meaning"],
      [["Base toe-out", f"{to['base_deg']} deg", "toe-out for a balanced hip"],
       ["Toe-out gain", f"{to['k']} deg per deg", "extra toe-out per degree of ER dominance"],
       ["Toe-out limits", f"{to['min_deg']} to {to['max_deg']} deg", "clipping range"],
       ["Base width", f"{w['base_factor']} x hip width", "width for a balanced hip"],
       ["Width gain", f"{w['k']} per deg", "extra width per degree of ER dominance"],
       ["Width limits", f"{w['min_factor']} to {w['max_factor']} x hip width", "clipping range"],
       ["Reference arc", f"{RULES['arc_reference_deg']} deg", "arcs below this damp the width change"],
       ["Asymmetry tolerance", f"{RULES['symmetry']['tol_deg']} deg", "left-right bias difference that flags asymmetry"],
       ["Pattern dead zone", "5 deg", "|bias| within this is 'balanced'"]],
      [4.0, 4.2, 6.3], left_cols=(2,))

H2("3.5 Squat Check: Rule-Based Biomechanical Classifier")
H3("3.5.1 Stand-up gate")
P("Right after the seated stage the knees are bent, which a counter would read as the bottom of a squat, so "
  "counting starts only after the average knee angle has stayed at 145 degrees or more for one second. A first gate "
  "of 155 degrees locked out people whose straight standing knee MediaPipe reads at 151 to 154 degrees (Section "
  "5.3). If the hip line is turned more than 20 degrees from square-on, the app asks the person to face the camera, "
  "because the stance measurements need a front view.")
H3("3.5.2 Per-frame signals")
P("For every frame the classifier computes the left and right knee angles (hip-knee-ankle), the height of the hip "
  "midpoint above the ankle midpoint in femur lengths, a depth ratio, the trunk lean (the hip-to-shoulder line "
  "against the vertical), the stance and knee widths on the ground plane in hip widths, the left-right difference in "
  "hip drop, the knee-over-foot offset and the facing angle. The depth ratio is")
equation("depth = (y(hip) − y(knee)) / femur length + 1", "3.8")
P("where y(hip) and y(knee) are the heights of the hip and knee midpoints (y points down). It is 0 when standing "
  "and 1 when the hips are level with the knees, so values above 1 mean below parallel.")
H3("3.5.3 Smoothing and velocity")
P("MediaPipe landmarks jitter, and now and then a single frame is badly wrong. Each signal goes through a moving "
  "median over the last 0.3 seconds, which removes single-frame spikes, and then exponential smoothing with a 0.12-second "
  "time constant; the vertical hip velocity is taken over the last 0.2 seconds. Frames with mean leg visibility below "
  "0.5 are skipped. Everything runs on timestamps, so a 30 fps webcam and a 15 fps recording give the same result.")
H3("3.5.4 State machine")
P("Figure 3.4 shows the cycle the classifier follows. A cycle can start only from upright standing, and it is "
  "scored only after the person has descended, turned at the bottom, risen and come back close to standing. Any "
  "phase that lasts longer than 4 seconds aborts the cycle.")
figure("fig_state_machine.png", 11.0, "State machine of the squat classifier with its transition conditions "
       "(thresholds as implemented).", "3.4")
H3("3.5.5 Biomechanical squat score")
P("A closed cycle is scored against the nine checks in Table 3.3, one point for each check passed. A score of at "
  "least 7 out of 9 makes it a squat repetition, so no single feature can accept or reject a movement on its own. "
  "Development testing added two refinements. The standing-return check is judged over the 0.5 seconds after the "
  "cycle closes, because people are still finishing the rise at that moment. Cycles whose geometry no body can "
  "produce (trunk beyond 100 degrees, a hip drop above 1.5 femur lengths or a depth ratio above 1.6) are treated as "
  "tracking noise and never accepted, whatever their score. A video counts as a squat set when at least two "
  "repetitions are accepted.")
table("3.3", "The nine checks of the biomechanical squat score",
      ["Check", "Passes when"],
      [["Depth", "maximum depth ratio >= 0.45"],
       ["Knee flexion", "knee bends >= 40 deg from standing and the bottom angle is <= 130 deg"],
       ["Hip motion", "hip drops >= 0.35 femur lengths and knee angle and hip height correlate >= 0.6"],
       ["Descent", "lasts 0.12-3.5 s and at least 60% of its frames move downward"],
       ["Bottom", "a turning point is found and the bottom lasts <= 2.5 s"],
       ["Ascent", "lasts 0.12-3.5 s and at least 60% of its frames move upward"],
       ["Standing return", "within 0.5 s the knee is back within 15 deg of standing and the trunk <= 40 deg"],
       ["Bilateral", "the smaller knee bend is >= 60% of the larger and mean hip-drop asymmetry <= 0.35 femur"],
       ["Upright trunk", "trunk <= 40 deg when standing and <= 70 deg throughout the cycle"]],
      [3.5, 11.0], left_cols=(1,))
figure("fig_signal_trace.png", 11.5, "The classifier on a real front-view squat. Shading shows the detected phases "
       "and each repetition is labelled with its score. The first dip, before 2 s, is not scored because the person "
       "had not yet been seen standing.", "3.5")
P("Every completed cycle is counted, but only accepted cycles are squat repetitions. A rejected cycle carries the "
  "checks it failed, which the user sees; one rejected for depth and hip motion, for example, was a partial squat.")
H2("3.6 Squat-Based Stance Analysis")
P("The confirmed repetitions are summarised by their medians: stance width at the ankles, depth ratio, maximum trunk "
  "lean and the knee-to-ankle width ratio at the bottom. The measured stance is NARROW below 1.4 hip widths, WIDE "
  "above 2.1 and MODERATE in between. The suggestion changes the stance only when a measured problem points at the "
  "width itself:")
bullets(["One step wider when the depth stays well above parallel (depth ratio below 0.5) and the trunk folds forward "
         "past 45 degrees at the same time, a sign that the hips have no room to sit between the feet. Either sign "
         "alone is not a width problem: a deep squat needs some lean, and a shallow upright squat is a depth problem.",
         "One step narrower from WIDE when the knees fall inside the feet at the bottom (knee width below 0.75 of "
         "ankle width), a sign that the stance is wider than the hips can control.",
         "Otherwise the measured stance is kept, with a cue to push the knees out if they cave in."])
P("Advice is given only for squats filmed from the front (hip line within 20 degrees of square-on); from the side, "
  "stance and knee width lie along the camera's depth axis, which MediaPipe measures poorly.")
H2("3.7 Storage, Dashboard and Deployment")
P("Each assessment is saved in a local SQLite database: a session row, the hip profile, the starting stance, one row "
  "per counted repetition and a squat-run row with the squat-based suggestion and its measurements. The Streamlit "
  "dashboard has four tabs: Assess starts the camera window and shows the result, History lists every stored session "
  "and plots the profile over time, Verification shows the saved evaluation results, and Offline demo replays an "
  "MM-Fit recording when no camera is available. The camera window can also save an annotated video.")
P("A Dockerfile and a Compose file package the dashboard, tests and offline scripts, with data mounted from the "
  "host; since Docker Desktop on Windows cannot pass a webcam into a Linux container, the camera assessment runs on "
  "the host and writes to the shared database. The project also installs as a Windows desktop app.")

# ===================================================================== CHAPTER 4
chapter(4, "Experimental Setup")
H2("4.1 Dataset")
H3("4.1.1 Multi-View Fitness Video Dataset")
P("The main dataset is the Multi-View raw video dataset of seven fitness exercises with good and bad form labels "
  "(Prashanth et al., 2026), published on Mendeley Data under a CC BY 4.0 licence. Twenty-six people were filmed with "
  "smartphones in a gym from the front, side and diagonal, each doing seven exercises with good and bad form: "
  "bodyweight squat, standing dumbbell side bend (labelled abs), bent-over dumbbell row (back), alternating dumbbell "
  "bicep curl, push-up, standing dumbbell shoulder press and overhead dumbbell triceps extension. Every file name "
  "carries the participant, exercise, form and view. We ran MediaPipe once over every video and cached each clip's "
  "world landmarks, visibility scores and body-scale constants. The videos are 15 frames per second and typically "
  "about 30 seconds long.")
P("Before any evaluation we checked the archive for duplicates. Its 1,042 videos hold only 1,008 distinct "
  "recordings, because 33 recordings appear more than once: ten under two different participants, one as both a squat "
  "and a back exercise, and five under two camera views. Since the rules are set on some participants and tested on "
  "others, such copies could put an evaluation participant's video into the development data, and a video filed as "
  "both squat and back has no single correct label. We dropped every copy whose participant, squat label or view was "
  "ambiguous and kept one copy otherwise, which removed 47 files. A re-check found no recording left under two people "
  "and none with conflicting squat labels. Participant 25 has no squat videos, and that person's four triceps videos "
  "stay in as extra non-squat test data.")
ex_counts = Counter(r["exercise"] for f in ("dev", "eval") for r in csv.DictReader(open(f"models/rule_classifier/results_{f}.csv")))
names = {"abs": "Standing dumbbell side bend (abs)", "back": "Bent-over dumbbell row (back)",
         "bicep_curl": "Alternating bicep curl", "push_up": "Push-up", "shoulder": "Standing shoulder press",
         "squat": "Bodyweight squat", "tricep": "Overhead triceps extension"}
rows = [[names[e], ex_counts[e], "squat" if e == "squat" else "non-squat"] for e in sorted(ex_counts)]
rows.append(["Total", sum(ex_counts.values()), f"{ex_counts['squat']} squat / {sum(ex_counts.values()) - ex_counts['squat']} non-squat"])
table("4.1", "Videos per exercise after removing duplicate recordings (all three views, 26 people)",
      ["Exercise", "Videos", "Class for the squat check"], rows, [7.0, 2.5, 5.0])
H3("4.1.2 Development and evaluation split for the rule classifier")
P("The rule classifier has no trained parameters, but its thresholds were tuned while looking at data, so that data "
  "cannot also be used to test it. Before any rule was written, the 26 participants were shuffled once with a fixed "
  "seed and split: 6 development participants, on whom the rules could change, and 20 evaluation participants, who "
  "were scored exactly once with the rules frozen (Table 4.2). The frozen configuration was saved with its hash "
  "before the evaluation run.")
table("4.2", "Person-level split used for the rule classifier",
      ["Group", "Participants", "Videos", "Squat / non-squat"],
      [["Development", ", ".join(map(str, RC["development_people"])), RCD["metrics"]["videos"],
        f"{RCD['metrics']['squat']} / {RCD['metrics']['non_squat']}"],
       ["Evaluation", ", ".join(map(str, RC["evaluation_people"])), M["videos"], f"{M['squat']} / {M['non_squat']}"]],
      [2.6, 7.2, 1.8, 2.9], size=10)
P(f"The classes are unbalanced. {pc(MAJORITY)}% of all six-second stretches of video in the cleaned dataset show an "
  f"exercise other than a squat, so a rule that never accepted a squat would still be right {pc(MAJORITY)}% of the "
  f"time. That is why Chapter 5 reports recall, specificity, F1 and balanced accuracy beside plain accuracy.")
H3("4.1.3 MM-Fit")
P("MM-Fit (Strömbäck et al., 2020) has 21 workout sessions by 10 people covering ten exercises, including 64 labelled "
  "squat sets, with 2D and 3D pose from its own estimator. Its labels mark whole exercise sets rather than single "
  "repetitions, and 61 of the 64 squat sets contain exactly ten repetitions, so a constant guess of ten is right "
  "95.3% of the time and the dataset cannot grade a repetition counter. In this report MM-Fit is used only for the "
  "offline demonstration (Section 5.3), as recordings the rules never saw, made with a different pose estimator.")
H2("4.2 Evaluation Metrics")
P("The squat check is a binary decision, so results come from the confusion counts: true positives (squat videos "
  "accepted), false negatives (squat videos missed), false positives (non-squat videos accepted) and true negatives. "
  "Squats are the minority class, so accuracy alone is misleading, and the balance-aware metrics are reported beside "
  "it (Brodersen et al., 2010).")
equation("Accuracy = (TP + TN) / (TP + TN + FP + FN)", "4.1")
equation("Precision = TP / (TP + FP),   Recall = TP / (TP + FN)", "4.2")
equation("F1 = 2 · Precision · Recall / (Precision + Recall)", "4.3")
equation("Specificity = TN / (TN + FP)", "4.4")
equation("Balanced accuracy = (Recall + Specificity) / 2", "4.5")
equation("FPR = FP / (FP + TN),   FNR = FN / (FN + TP)", "4.6")
P("Scoring is per video: a video counts as a squat set when at least two of its repetitions are accepted. "
  "Repetitions wrongly accepted in non-squat videos are counted as well, because a user would see them as false "
  "squats. Neither dataset labels single repetitions, so repetition counts are compared with a reference count of "
  "knee-angle valleys in the same video.")
H2("4.3 Implementation Details")
P(f"Development and testing ran on a Windows 11 laptop using only the CPU, with code written in Python 3.11 (Table "
  f"4.3). The person split uses a fixed seed and the frozen rule configuration is stored with its hash, so the "
  f"evaluation can be repeated exactly. The project has {N_TESTS} automated tests covering the rotation geometry and "
  f"its sign convention, the controller's stages, the rule classifier on synthetic squats and non-squats, frame-rate "
  f"independence, the storage layer and the dashboard. All of them pass, and the Docker image runs the same suite.")
table("4.3", "Software used", ["Component", "Version", "Purpose"],
      [["Python", "3.11", "implementation language"],
       ["MediaPipe", "0.10.14", "BlazePose pose estimation"],
       ["OpenCV", "4.11", "camera capture, drawing, video I/O"],
       ["NumPy / SciPy", "1.26.4 / 1.17", "signal processing and statistics"],
       ["Streamlit", "1.37.1", "dashboard"],
       ["SQLite", "Python standard library", "assessment storage"],
       ["Matplotlib", "3.11", "figures"],
       ["pytest", "9.1", "automated tests"],
       ["Docker / Compose", "29.5 / 5.1", "containerised dashboard and tools"]],
      [4.2, 4.3, 6.0])

# ===================================================================== CHAPTER 5
chapter(5, "Results & Discussion")
H2("5.1 Rule-Based Squat Classifier on Held-Out People")
P(f"The rules went through three versions, each tested only on the six development participants. The first found "
  f"28 of the 32 development squat videos with no false alarms; it missed people who squat without pausing at the "
  f"top, a bottom turn too quick for 15 frames per second, and a standing-return check judged before the person had "
  f"finished rising. After those fixes and the noise gate in Section 3.5.5, the development set was classified at "
  f"{pc(RCD['metrics']['accuracy'])}% accuracy ({RCD['metrics']['TP']} of {RCD['metrics']['squat']} squat videos, "
  f"{RCD['metrics']['FP']} false alarm in {RCD['metrics']['non_squat']}). The rules were then frozen and the "
  f"{len(RC['evaluation_people'])} evaluation participants were scored once.")
fr = [r for r in EVAL_ROWS if r["view"] == "front"]
def counts(rows):
    tp = sum(r["ground_truth"] == "squat" and r["prediction"] == "squat" for r in rows)
    fn = sum(r["ground_truth"] == "squat" and r["prediction"] != "squat" for r in rows)
    fp = sum(r["ground_truth"] != "squat" and r["prediction"] == "squat" for r in rows)
    tn = sum(r["ground_truth"] != "squat" and r["prediction"] != "squat" for r in rows)
    pr, rc, sp = tp / (tp + fp), tp / (tp + fn), tn / (tn + fp)
    return dict(n=len(rows), TP=tp, FN=fn, FP=fp, TN=tn, acc=(tp + tn) / len(rows), prec=pr, rec=rc,
                f1=2 * pr * rc / (pr + rc), spec=sp, bal=(rc + sp) / 2, fpr=fp / (fp + tn), fnr=fn / (fn + tp))
F = counts(fr)
table("5.1", "Rule-based squat classifier on the 20 evaluation participants (scored once, rules frozen)",
      ["Metric", "All three views", "Front camera only"],
      [["Videos (squat / non-squat)", f"{M['videos']} ({M['squat']} / {M['non_squat']})",
        f"{F['n']} ({F['TP'] + F['FN']} / {F['FP'] + F['TN']})"],
       ["TP / FN / FP / TN", f"{M['TP']} / {M['FN']} / {M['FP']} / {M['TN']}", f"{F['TP']} / {F['FN']} / {F['FP']} / {F['TN']}"],
       ["Accuracy", f"{pc(M['accuracy'])}%", f"{pc(F['acc'])}%"],
       ["Precision", f"{pc(M['precision'])}%", f"{pc(F['prec'])}%"],
       ["Recall", f"{pc(M['recall'])}%", f"{pc(F['rec'])}%"],
       ["F1-score", f"{pc(M['f1'])}%", f"{pc(F['f1'])}%"],
       ["Specificity", f"{pc(M['specificity'])}%", f"{pc(F['spec'])}%"],
       ["Balanced accuracy", f"{pc(M['balanced_accuracy'])}%", f"{pc(F['bal'])}%"],
       ["False positive rate", f"{pc(M['false_positive_rate'])}%", f"{pc(F['fpr'])}%"],
       ["False negative rate", f"{pc(M['false_negative_rate'])}%", f"{pc(F['fnr'])}%"]],
      [5.0, 4.8, 4.7])
figure("fig_confusion_rules.png", 7.0, "Confusion matrix of the rule-based classifier on the evaluation participants.", "5.1")
ns = [r for r in EVAL_ROWS if r["ground_truth"] == "non-squat"]
sq = sorted(int(r["rep_count"]) for r in EVAL_ROWS if r["ground_truth"] == "squat")
P(f"Non-squat exercises are rejected reliably: {pc(M['specificity'])}% of non-squat videos were rejected. "
  f"Relative to class size, missed squats are the larger error, at a {pc(M['false_negative_rate'])}% false negative "
  f"rate. At the level of single repetitions, {sum(int(r['rep_count']) for r in ns)} repetitions were wrongly "
  f"accepted across {sum(int(r['rep_count']) > 0 for r in ns)} of the {len(ns)} non-squat videos, while a squat "
  f"video gave a median of {sq[len(sq) // 2]} accepted repetitions.")
figure("fig_per_exercise.png", 13.5, "Correct decisions by exercise on the 20 evaluation participants. Green: squat "
       "videos found (recall); blue: videos of each other exercise correctly rejected. Labels give the count correct "
       "out of the total, and the white text the errors. The dashed line is the overall accuracy; the vertical axis "
       "starts at 80%.", "5.2")
# per-exercise reading of Figure 5.2, counted from the evaluation results
_nm = {"abs": "side bend", "back": "bent-over row", "bicep_curl": "bicep curl", "push_up": "push-up",
       "shoulder": "shoulder press", "tricep": "triceps extension"}
_cnt = {}
for _r in EVAL_ROWS:
    _c = _cnt.setdefault(_r["exercise"], [0, 0]); _c[0] += 1; _c[1] += int(_r["correct"])
_fa = sorted(((n - k, e) for e, (n, k) in _cnt.items() if e != "squat" and n > k), key=lambda t: (-t[0], t[1]))
_clean = [_nm[e] + "s" for e, (n, k) in sorted(_cnt.items()) if e != "squat" and n == k]
_word = ["no", "one", "two", "three", "four", "five", "six"]
_and = lambda xs: xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]
_sq_n, _sq_ok = _cnt["squat"]
_worst_n, _worst_ok = _cnt[_fa[0][1]]
P(f"Figure 5.2 splits the decisions by exercise. The green bar is recall: {_sq_ok} of {_sq_n} squat videos were "
  f"found and {_sq_n - _sq_ok} were missed. Each blue bar is the share of that exercise's videos correctly rejected. "
  f"The {M['FP']} false alarms come from {_word[len(_fa)]} exercises: the {_nm[_fa[0][1]]} gave {_fa[0][0]} of them "
  f"({_fa[0][0]} of its {_worst_n} videos), and the {_and([_nm[e] for _, e in _fa[1:]])} one each. "
  f"{_and(_clean).capitalize()} were never accepted, since the person is either not upright or does not bend the "
  f"knees. Even the {_nm[_fa[0][1]]}, the hardest exercise to reject, was rejected in {pc(_worst_ok / _worst_n)}% of "
  f"its videos, which is why the chart's axis starts at 80%: on a full 0 to 100% scale these differences would "
  f"not be visible.")
H2("5.2 Error Analysis")
P(f"All {M['FN'] + M['FP']} wrong decisions on the evaluation participants were examined with the rules still "
  f"frozen. They fall into four groups (Table 5.2).")
table("5.2", "Wrongly classified evaluation videos, grouped by cause",
      ["Cause", "Videos", "What the measurements show"],
      [["Shallow 'bad form' squats", "5 missed", "most cycles fail the depth, knee-flexion and hip-motion checks together "
        "(person 8: knee bends of 27-29 deg); by the rules these are partial movements, but the dataset labels them as squats"],
       ["Tracking failures", "3 missed", "a side-view clip where the knee never reads below 148 deg, a clip with trunk "
        "readings up to 158 deg, and a side-view clip whose trunk and left-right readings fail in 5 of 6 cycles"],
       ["Too few complete cycles", "1 missed", "only one complete cycle detected; the video rule needs two"],
       ["Knee bends in other exercises", "6 false alarms", "a shoulder press with real knee bends down to 78 deg, three "
        "bent-over rows, one side bend and one bicep curl, scoring 7 to 9 of 9"]],
      [4.0, 2.4, 8.1], size=10, left_cols=(2,))
P("The two largest groups have different causes. The shallow squats are a disagreement between the rules and the "
  "labels: the dataset calls a quarter-depth movement a (bad) squat, while our rules require reasonable depth. The "
  "false alarms come from exercises in which the person really does bend the knees and lower the hips, which a "
  "classifier that looks only at the legs and trunk cannot always separate from a squat. Many of those cycles scored "
  "exactly 7, the threshold, and the check they failed most often was left-right consistency.")
H2("5.3 Real-Time Pipeline Tests")
e15 = [r for r in E2E if r["stream_fps"] == 15.0]
rows = []
for r in e15:
    rows.append([r["video"].replace("subject_", "s").replace("_", " "), r["truth"], r["prediction"],
                 "yes" if r["correct"] else "no", f"{r['reps_confirmed']} / {r['reps_counted']}",
                 r["reference_reps"] if r["reference_reps"] is not None else "-"])
_sq_ok = sum(r["truth"] == "squat" and r["correct"] for r in e15)
P(f"The offline evaluation reads cached landmarks. To test the live path, raw videos of evaluation participants were "
  f"decoded with OpenCV, run through the application's own MediaPipe pose estimator and fed frame by frame, with "
  f"timestamps, into the same squat runner the live application uses. {sum(r['correct'] for r in e15)} of "
  f"{len(e15)} videos were classified correctly (Table 5.3). For the {_sq_ok} correctly detected squat videos, the "
  f"number of accepted repetitions was within one of the reference count. The one error is a shallow squat video of "
  f"the kind described in Section 5.2.")
table("5.3", "Raw video through the live chain (MediaPipe + rule classifier, 15 fps)",
      ["Video", "Truth", "Prediction", "Correct", "Accepted / cycles", "Reference reps"], rows,
      [5.0, 1.8, 1.9, 1.4, 2.4, 2.0], size=9.5)
e30 = {r["video"]: r for r in E2E if r["stream_fps"] == 30.0}
same = [v for v, r in e30.items() for x in e15 if x["video"] == v and x["reps_confirmed"] == r["reps_confirmed"]
        and x["prediction"] == r["prediction"]]
P(f"Webcams usually deliver 30 frames per second. Replaying the videos as a 30 fps stream (each frame shown twice at "
  f"1/30-second intervals) gave the same decision and repetition count as the native 15 fps replay for {len(same)} "
  f"of {len(e30)} videos tested, so the timestamp-based design does not depend on the frame rate.")
P("The full controller was also driven end to end with real recorded squats, from calibration to saving. A "
  "good-form squat video gave five confirmed repetitions, each scoring 9 of 9, and the set ended by itself; a bad-form "
  "video gave four confirmed and two rejected repetitions, the rejected ones having impossible geometry from tracking "
  "noise; a bicep-curl video gave no repetitions and no stance suggestion. These tests exposed the stand-up problem "
  "from Section 3.5.1: at 155 degrees three of 25 people could never start, while at 145 degrees each person's own "
  "upright frames held for 1.5 seconds passed the gate in all 50 front-view squat videos.")
_t19, _t05 = DEMO_TOTAL["w19"], DEMO_TOTAL["w05"]
P(f"Finally, the offline demonstration replays MM-Fit workouts, which the rules never saw and which come from a "
  f"different pose estimator. On workout w19 the classifier accepted {_t19[2]} of the {_t19[0]} labelled squat "
  f"repetitions, and on workout w05 it counted and accepted {_t05[2]} of {_t05[0]}.")
H2("5.4 Squat Analysis and Stance Suggestions")
fs = [r for r in fr if r["stance"]]
cur = Counter(r["stance"] for r in fs); rec = Counter(r["recommendation"] for r in fs)
changed = sum(r["stance"] != r["recommendation"] for r in fs)
P(f"Every squat video with confirmed repetitions gets a stance suggestion from its measurements. Figure 5.3 compares "
  f"the measured and suggested stances for the {len(fs)} front-view evaluation videos. The suggestion differed from "
  f"the measured stance in {changed} of them; in the rest no width-related problem was measured and the advice was to "
  f"keep the stance. An earlier version of the rule widened the stance whenever depth or trunk lean alone looked "
  f"poor, and it recommended WIDE for 67 of 100 evaluation squat videos. Requiring both signs together removed that "
  f"skew.")
figure("fig_stance_front.png", 9.5, "Measured and suggested stance for front-view evaluation squats.", "5.3")
P("The comparison also exposed a front-camera bias. For the development participants, front-view squats gave a "
  "median depth ratio of 0.45 against 0.72 from the side, and a median trunk lean of 45 degrees against 26. The front "
  "view is still right for stance width and knee tracking, which lie across the image, but the depth and lean "
  "thresholds have to be read with this bias in mind. No dataset we had carries stance labels, so the suggestions "
  "show that the rule behaves sensibly, not that it is correct.")
H2("5.5 Hip-Rotation Profiling")
P("The hip-rotation stage was checked for correct computation rather than accuracy. Unit tests cover the shank-angle "
  "geometry against known angles, the sign convention on both legs, the neutral correction, the hold-to-lock timing "
  "and the profile and stance formulas. The controller test above drives the whole seated protocol with synthetic "
  "seated frames scaled to a real person's femur and hip width, because no public dataset contains the seated "
  "rotation movement.")
P(f"Section 5.6 follows one real assessment through the live application. The offline MM-Fit demonstration applies "
  f"the same geometry to unconstrained workout footage, where the total arc reaches {DEMO_ARC:.0f} degrees (workout "
  f"{DEMO_WID}), beyond any plausible hip range, and the application warns that the reading is dominated by "
  f"whole-body movement rather than hip rotation. That is intended, because the reading means something only in the "
  f"guided seated position. Validating it against a goniometer is left for future work.")
H2("5.6 A Live Assessment")
_lp, _ls, _lsum, _lsq = LIVE_P, LIVE_S, LIVE_SUM, LIVE_SQ
P(f"Figures 5.4 and 5.5 are frames from one assessment recorded with the application on {LIVE_DAY} (session "
  f"#{LIVE_ID}), and Figures 5.6 and 5.7 show its result in the dashboard. The camera window shows the instruction at "
  "the top, the live reading at the bottom left and a progress bar while a reading is held. Each rotation step opens "
  "with a four-second countdown (Figure 5.4b); on the frame where a reading locks, the header already shows the next "
  "step.")
P(f"The left leg locked at {_lp['ir_left']:.0f} degrees of internal and {_lp['er_left']:.0f} degrees of external "
  f"rotation, and the right leg at {_lp['ir_right']:.0f} and {_lp['er_right']:.0f} degrees (Table 5.4). Averaged over "
  f"both legs that is {_lp['ir_max']:.1f} degrees of internal and {_lp['er_max']:.1f} degrees of external rotation, a "
  f"total arc of {_lp['total_arc']:.1f} degrees and a bias of {_lp['rotation_bias']:+.1f} degrees, which the profiler "
  f"classes as {_lp['pattern']}. The two legs' biases differ by {_lp['symmetry']:.1f} degrees, more than the "
  f"{RULES['symmetry']['tol_deg']:.0f}-degree tolerance, so the starting stance is flagged as asymmetric; a second "
  "session would show whether a gap this large comes from the hips or from how far each leg was turned. The rules "
  f"suggested starting at {_ls['width_factor']:.2f} times hip width "
  f"with {_ls['toe_out_deg']:.0f} degrees of toe-out, narrower and straighter than the generic default, as an "
  f"{_lp['pattern']} profile calls for.")
figure("fig_live_seated.png", 14.5, "Seated stage of a live assessment, from the application's own recording. "
       "Positive readings are internal rotation and negative ones external rotation; the locked values are the ones "
       "stored in the hip profile.", "5.4")
figure("fig_live_squat.png", 11.5, "Standing stage of the same assessment: the stand-up gate, a squat in progress "
       "and confirmed repetitions (the long confirmation banner runs past the edge of the frame).", "5.5")
P(f"After the seated stage the person stood up, and counting began once both knees read upright for one second "
  f"(Figure 5.5a). All {_lsum['counted']} repetitions passed the nine-check rules and were confirmed as squats. Their "
  f"depth ratio averaged {_lsum['mean_depth']:.2f}, with a best of {_lsum['max_depth']:.2f}, where 1.0 means the "
  f"thighs are parallel to the floor. From the confirmed squats the system measured a stance of "
  f"{_lsq['stance_width']:.2f} hip widths ({_lsq['current_stance']}), a maximum trunk lean of "
  f"{_lsq['max_trunk_lean_deg']:.0f} degrees and a hip line {_lsq['facing_deg']:.0f} degrees from square-on, inside "
  f"the {FRONT_FACING_MAX_DEG:.0f}-degree limit for a front view. No width-related problem was found, so the advice "
  f"was to keep the {_lsq['recommendation']} stance.")
P("The two suggestions answer different questions. The seated test gives a starting point before any squat, while "
  "the squat-based advice judges the stance the person actually used. Here the person squatted much wider than the "
  "starting suggestion and had no problem at that width, so the squat-based advice asked for no change.")
_deg = "°"
table("5.4", f"The live assessment of {LIVE_DAY}, read from the application's database (session #{LIVE_ID})",
      ["Quantity", "Value"],
      [["Left leg: internal / external rotation", f"{_lp['ir_left']:.1f}{_deg} / {_lp['er_left']:.1f}{_deg}"],
       ["Right leg: internal / external rotation", f"{_lp['ir_right']:.1f}{_deg} / {_lp['er_right']:.1f}{_deg}"],
       ["Mean internal / external rotation", f"{_lp['ir_max']:.1f}{_deg} / {_lp['er_max']:.1f}{_deg}"],
       ["Total arc, bias, pattern", f"{_lp['total_arc']:.1f}{_deg}, {_lp['rotation_bias']:+.1f}{_deg}, {_lp['pattern']}"],
       ["Left/right bias difference", f"{_lp['symmetry']:.1f}{_deg}" + (" (asymmetric)" if _ls["asymmetric"] else "")],
       ["Starting stance from the rules", f"{_ls['width_factor']:.2f} x hip width, {_ls['toe_out_deg']:.0f}{_deg} toe-out"],
       ["Repetitions counted / confirmed", f"{_lsum['counted']} / {_lsum['confirmed']}"],
       ["Depth ratio, mean / best", f"{_lsum['mean_depth']:.2f} / {_lsum['max_depth']:.2f}"],
       ["Measured squat stance", f"{_lsq['current_stance']}, {_lsq['stance_width']:.2f} x hip width"],
       ["Squat-based suggestion", f"{_lsq['recommendation']}"]],
      [7.5, 6.5])
figure("dash_result_profile.png", 11.0, "Dashboard result for the same session: the hip profile, the reading for each "
       "leg and the starting stance.", "5.6")
figure("dash_result_squat.png", 11.0, "Dashboard result for the same session: the squat check, one row per "
       "repetition, and the squat-based stance advice.", "5.7")
H2("5.7 The Application")
P("Figure 5.8 shows the History tab, which lists every stored assessment and opens on the most recent one (here the "
  "session from Section 5.6). Figure 5.9 shows the Verification tab, which reads the classifier's frozen evaluation "
  "from the saved result files. The same dashboard runs from the Docker image and, on Windows, as a desktop app in "
  "its own window.")
figure("dashboard_history.png", 11.5, "Dashboard, History tab: a stored assessment with its hip profile and starting "
       "stance.", "5.8")
figure("dashboard_verification.png", 11.5, "Dashboard, Verification tab: the rule classifier's evaluation on the "
       "held-out participants, read from the saved results.", "5.9")
H2("5.8 Discussion")
P(f"A classifier built only from explicit biomechanical rules, with no training data, reached an F1-score of "
  f"{pc(M['f1'])}% with {pc(M['specificity'])}% specificity when checked once on {len(RC['evaluation_people'])} people "
  f"it had never seen. Its errors make sense: missed squats are mostly partial movements or tracking failures, false "
  f"alarms are other exercises with real knee bends, and each rejected repetition comes with the checks it failed. "
  f"The data needed as much care as the rules, since duplicate recordings filed under different people could have "
  f"leaked evaluation videos into development.")
P("The main limitations are those stated in Section 1.5. The rotation readings have not been validated against a "
  "goniometer. The stance suggestions are rules rather than a model fitted to outcomes, because no dataset links "
  "stance to an outcome for each person. The evaluation used one dataset recorded under similar gym conditions, and "
  "the system has not yet been tried by users at home with their own webcams.")

# ===================================================================== CHAPTER 6
chapter(6, "Conclusion & Future Work")
H2("6.1 Conclusion")
P(f"StanceSense-RT shows that one webcam and a laptop CPU are enough for a guided, personalised squat assessment. A "
  f"seated rotation test gives a functional hip-rotation profile, explicit rules turn it into a starting stance, and "
  f"a rule-based biomechanical classifier then confirms, counts and measures each squat, so the final stance "
  f"suggestion rests on measurements the person can see. On a public multi-view dataset, with duplicate recordings "
  f"removed and the rules frozen on {len(RC['development_people'])} development participants, the classifier reached "
  f"{pc(M['accuracy'])}% accuracy, {pc(M['precision'])}% precision, {pc(M['recall'])}% recall and an F1-score of "
  f"{pc(M['f1'])}% on {M['videos']} videos from {len(RC['evaluation_people'])} new participants, and "
  f"{pc(FRONT_ACC)}% accuracy on front-camera videos. Because the decision comes from rules, it needs no training "
  f"data, which matters when squats are a small minority of the labelled movement, and it can always say why a "
  f"repetition was accepted or rejected. The application stores every assessment, shows it in a dashboard and runs "
  f"locally, in Docker or as a desktop app.")
H2("6.2 Future Work")
bullets(["Validate the rotation readings against goniometer measurements taken by a physiotherapist on the same "
         "people, and report the agreement before making any angular claim.",
         "Collect squats from users at home with their own webcams, with stance and comfort recorded, to check the "
         "stance rules against real outcomes.",
         "Recalibrate the depth and trunk-lean thresholds on front-camera recordings, given the bias measured in "
         "Section 5.4.",
         "Add arm and upper-body features so that knee-bending exercises such as presses and rows are rejected more "
         "reliably, and name the movement when a repetition is rejected.",
         "Request access to datasets with expert error labels, such as Fitness-AQA, to extend the checks to named "
         "faults like the knees caving in.",
         "Give live audio or on-screen cues during the squat, based on the same per-repetition checks."])

# ===================================================================== REFERENCES
doc.add_heading("REFERENCES", level=1)
REFS = [
    "Bazarevsky, V., Grishchenko, I., Raveendran, K., Zhu, T., Zhang, F., & Grundmann, M. (2020). BlazePose: On-device "
    "real-time body pose tracking (arXiv:2006.10204). arXiv. https://doi.org/10.48550/arXiv.2006.10204",
    "Brodersen, K. H., Ong, C. S., Stephan, K. E., & Buhmann, J. M. (2010). The balanced accuracy and its posterior "
    "distribution. In 2010 20th International Conference on Pattern Recognition (pp. 3121–3124). IEEE. "
    "https://doi.org/10.1109/ICPR.2010.764",
    "Cao, Z., Hidalgo, G., Simon, T., Wei, S.-E., & Sheikh, Y. (2021). OpenPose: Realtime multi-person 2D pose "
    "estimation using part affinity fields. IEEE Transactions on Pattern Analysis and Machine Intelligence, 43(1), "
    "172–186. https://doi.org/10.1109/TPAMI.2019.2929257",
    "Capecci, M., Ceravolo, M. G., Ferracuti, F., Iarlori, S., Monteriù, A., Romeo, L., & Verdini, F. (2019). The "
    "KIMORE dataset: KInematic assessment of MOvement and clinical scores for remote monitoring of physical "
    "REhabilitation. IEEE Transactions on Neural Systems and Rehabilitation Engineering, 27(7), 1436–1448.",
    "Dwibedi, D., Aytar, Y., Tompson, J., Sermanet, P., & Zisserman, A. (2020). Counting out time: Class agnostic "
    "video repetition counting in the wild. In Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern "
    "Recognition (pp. 10387–10396).",
    "Escamilla, R. F. (2001). Knee biomechanics of the dynamic squat exercise. Medicine and Science in Sports and "
    "Exercise, 33(1), 127–141. https://doi.org/10.1097/00005768-200101000-00020",
    "Escamilla, R. F., Fleisig, G. S., Lowry, T. M., Barrentine, S. W., & Andrews, J. R. (2001). A three-dimensional "
    "biomechanical analysis of the squat during varying stance widths. Medicine and Science in Sports and Exercise, "
    "33(6), 984–998. https://doi.org/10.1097/00005768-200106000-00019",
    "Lugaresi, C., Tang, J., Nash, H., McClanahan, C., Uboweja, E., Hays, M., Zhang, F., Chang, C.-L., Yong, M. G., "
    "Lee, J., Chang, W.-T., Hua, W., Georg, M., & Grundmann, M. (2019). MediaPipe: A framework for building perception "
    "pipelines (arXiv:1906.08172). arXiv. https://doi.org/10.48550/arXiv.1906.08172",
    "Parmar, P., Gharat, A., & Rhodin, H. (2022). Domain knowledge-informed self-supervised representations for "
    "workout form assessment. In Computer Vision – ECCV 2022. Lecture Notes in Computer Science (Vol. 13698). "
    "Springer.",
    "Prashanth, D. S., Kamath, P., Bhandary, R., Pai, S. A., PN, P., P, A., & Mashetty, A. (2026). A multi-view raw "
    "video dataset of seven fitness exercises with good/bad form labels (Version 2) [Data set]. Mendeley Data. "
    "https://doi.org/10.17632/kgbb3yn47p.2",
    "Roach, K. E., & Miles, T. P. (1991). Normal hip and knee active range of motion: The relationship to age. "
    "Physical Therapy, 71(9), 656–665. https://doi.org/10.1093/ptj/71.9.656",
    "Ruwe, P. A., Gage, J. R., Ozonoff, M. B., & DeLuca, P. A. (1992). Clinical determination of femoral anteversion: "
    "A comparison with established techniques. The Journal of Bone and Joint Surgery. American Volume, 74(6), "
    "820–830.",
    "Schoenfeld, B. J. (2010). Squatting kinematics and kinetics and their application to exercise performance. "
    "Journal of Strength and Conditioning Research, 24(12), 3497–3506. https://doi.org/10.1519/JSC.0b013e3181bac2d7",
    "Strömbäck, D., Huang, S., & Radu, V. (2020). MM-Fit: Multimodal deep learning for automatic exercise "
    "logging across sensing devices. Proceedings of the ACM on Interactive, Mobile, Wearable and Ubiquitous "
    "Technologies, 4(4), Article 168. https://doi.org/10.1145/3432701",
    "Vakanski, A., Jun, H.-P., Paul, D., & Baker, R. (2018). A data set of human body movements for physical "
    "rehabilitation exercises. Data, 3(1), Article 2. https://doi.org/10.3390/data3010002",
    "Zhao, Z., Kiciroglu, S., Vinzant, H., Cheng, Y., Katircioglu, I., Salzmann, M., & Fua, P. (2022). 3D pose based "
    "feedback for physical exercises. In Proceedings of the Asian Conference on Computer Vision (ACCV).",
]
for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(8)
    p.add_run(f"[{i}]. {ref}")

os.makedirs(OUTDIR, exist_ok=True)
doc.save(OUT)
print("saved", OUT, "| figures", len(FIGS), "| tables", len(TABS))
